import streamlit as st
import pandas as pd
import random
from docx import Document
from io import BytesIO

# -----------------------------
# CONFIG
# -----------------------------
st.set_page_config(page_title="UrbanLab", layout="wide")

st.markdown("""
<div style="background-color:#1e3a8a;color:white;padding:12px;border-radius:8px;text-align:center;font-weight:bold;margin-bottom:20px;">
🏙️ SISTEMA DE SIMULACIÓN DE POLÍTICA CRIMINAL
</div>
""", unsafe_allow_html=True)

st.title("UrbanLab — Teoría de la Desorganización Social")

# -----------------------------
# LISTA DE BARRIOS
# -----------------------------
if "barrios_disponibles" not in st.session_state:
    st.session_state.barrios_disponibles = [
        "Polígono Sur (Sevilla)",
        "El Raval (Barcelona)",
        "El Cabanyal (Valencia)",
        "Puente de Vallecas (Madrid)",
        "Usera (Madrid)",
        "Ciutat Meridiana (Barcelona)",
        "La Mina (Sant Adrià del Besòs)"
    ]

# -----------------------------
# IDENTIFICACIÓN DE GRUPO
# -----------------------------
st.header("👥 Identificación del grupo")

grupo_input = st.text_input("Introduce el nombre del grupo")

if "grupo" not in st.session_state:
    st.session_state.grupo = None

if "barrio" not in st.session_state:
    st.session_state.barrio = None

# asignación única
if grupo_input and st.session_state.grupo is None:

    st.session_state.grupo = grupo_input

    if len(st.session_state.barrios_disponibles) > 0:
        barrio = random.choice(st.session_state.barrios_disponibles)
        st.session_state.barrio = barrio
        st.session_state.barrios_disponibles.remove(barrio)
    else:
        st.session_state.barrio = "No quedan barrios disponibles"

# mostrar barrio
if st.session_state.barrio:

    st.header("📍 Barrio asignado")
    st.success(st.session_state.barrio)

    st.warning("⚠️ Este barrio es fijo y no puede modificarse")

# -----------------------------
# DIAGNÓSTICO
# -----------------------------
if st.session_state.barrio:

    st.header("Paso 1 — Diagnóstico criminológico")

    diagnostico = st.text_area("""
Analiza el barrio utilizando la teoría de la desorganización social:

- Cohesión social  
- Control social informal  
- Condiciones estructurales  
- Factores criminógenos  
""", height=200)

# -----------------------------
# PLAN
# -----------------------------
if st.session_state.barrio:

    st.header("Paso 2 — Diseño de intervención")

    plan = st.text_area("Describe tu plan de intervención", height=200)

# -----------------------------
# INTERPRETADOR
# -----------------------------
def interpretar_plan(plan):

    texto = plan.lower()

    categorias = {
        "Control formal": {
            "kw":["polic","vigilancia","cámaras"],
            "impacto":{"policia":15,"control":10},
            "tipo":"punitiva"
        },
        "Cohesión social": {
            "kw":["vecin","comunit","asociaciones"],
            "impacto":{"cohesion":15,"control":10},
            "tipo":"estructural"
        },
        "Urbanismo": {
            "kw":["urban","espacio","rehabilitación"],
            "impacto":{"desorganizacion":-15},
            "tipo":"estructural"
        },
        "Intervención económica": {
            "kw":["empleo","educa","formación"],
            "impacto":{"pobreza":-15},
            "tipo":"estructural"
        }
    }

    cambios = {"policia":0,"cohesion":0,"control":0,"desorganizacion":0,"pobreza":0}
    contribuciones = []
    tipos = set()

    for nombre,data in categorias.items():
        for kw in data["kw"]:
            if kw in texto:
                tipos.add(data["tipo"])
                for k,v in data["impacto"].items():
                    cambios[k]+=v
                contribuciones.append(nombre)
                break

    tipo_final = "mixta" if len(tipos)>1 else list(tipos)[0] if tipos else "indefinida"
    score = len(contribuciones)*2

    return cambios, contribuciones, tipo_final, min(score,10)

# -----------------------------
# EJECUCIÓN
# -----------------------------
if st.session_state.barrio:

    if st.button("Ejecutar simulación"):

        if len(diagnostico) < 50:
            st.warning("El diagnóstico es demasiado breve")
            st.stop()

        if len(plan) < 50:
            st.warning("El plan es demasiado breve")
            st.stop()

        cambios, contribuciones, tipo, score = interpretar_plan(plan)

        base = {
            "desorganizacion":70,
            "cohesion":40,
            "control":30,
            "pobreza":60,
            "policia":40
        }

        for k in cambios:
            if k in base:
                base[k]+=cambios[k]

        delito = (
            base["desorganizacion"]*0.4 +
            base["pobreza"]*0.3 -
            base["cohesion"]*0.3 -
            base["control"]*0.2
        )

        st.session_state.resultado = delito
        st.session_state.contrib = contribuciones
        st.session_state.tipo = tipo
        st.session_state.score = score

# -----------------------------
# RESULTADOS
# -----------------------------
if "resultado" in st.session_state:

    st.header("Paso 3 — Resultados")

    c1,c2,c3 = st.columns(3)
    c1.metric("Nivel de delincuencia", round(st.session_state.resultado,2))
    c2.metric("Tipo de estrategia", st.session_state.tipo)
    c3.metric("Puntuación teórica", st.session_state.score)

    df = pd.DataFrame(st.session_state.contrib, columns=["Intervenciones detectadas"])
    st.dataframe(df)

    st.subheader("🧠 Feedback del diagnóstico")

    if "cohesion" not in diagnostico.lower():
        st.warning("Falta análisis de la cohesión social")

    if "control" not in diagnostico.lower():
        st.warning("Falta análisis del control informal")

    if "pobre" not in diagnostico.lower():
        st.warning("Faltan factores estructurales")

# -----------------------------
# INFORME WORD
# -----------------------------
if "resultado" in st.session_state:

    st.header("Paso 4 — Generar informe")

    if st.button("Generar informe"):

        doc = Document()

        doc.add_heading('INFORME DE POLÍTICA CRIMINAL', 1)

        doc.add_paragraph(f"Grupo: {st.session_state.grupo}")
        doc.add_paragraph(f"Barrio: {st.session_state.barrio}")

        doc.add_heading('Diagnóstico',2)
        doc.add_paragraph(diagnostico)

        doc.add_heading('Intervención',2)
        doc.add_paragraph(plan)

        doc.add_heading('Resultados',2)
        doc.add_paragraph(f"Nivel de delincuencia: {round(st.session_state.resultado,2)}")
        doc.add_paragraph(f"Tipo de estrategia: {st.session_state.tipo}")

        buffer = BytesIO()
        doc.save(buffer)
        buffer.seek(0)

        st.download_button(
            "📄 Descargar informe",
            buffer,
            file_name="informe_urbanlab.docx"
        )

# -----------------------------
# DRIVE + EVALUACIÓN
# -----------------------------
st.divider()
st.header("📄 Intercambio y evaluación")

grupo = st.selectbox("Selecciona tu grupo", ["Grupo A","Grupo B","Grupo C", "Grupo D", "Grupo E", " Grupo F", "Grupo G"])

links = {
    "Grupo A":{"upload":"https://drive.google.com/drive/folders/1milDBVgP5qossw_UXJbMhZLnwQ9vOy7t?usp=sharing","review":"https://drive.google.com/drive/folders/1lljOVHNvhmwbqopny9nE8fIjH-S6nVWx?usp=drive_link"},
    "Grupo B":{"upload":"https://drive.google.com/drive/folders/1lljOVHNvhmwbqopny9nE8fIjH-S6nVWx?usp=drive_link","review":"https://drive.google.com/drive/folders/1Sw2Fj6dridgzOVh-9Q89ThJ6cvPc-tXw?usp=sharing"},
    "Grupo C":{"upload":"https://drive.google.com/drive/folders/1Sw2Fj6dridgzOVh-9Q89ThJ6cvPc-tXw?usp=sharing","review":"https://drive.google.com/drive/folders/1milDBVgP5qossw_UXJbMhZLnwQ9vOy7t?usp=drive_link"},
    "Grupo D":{"https://drive.google.com/drive/folders/1XHkysQiaxAevSWVrkgZaFvE8QPNUfMMl?usp=sharing","review":"https://drive.google.com/drive/folders/1bxb--QlXr4oqN0DkXE6_5bWywhY3AQt0?usp=sharing"}, 
    "Grupo E": {"https://drive.google.com/drive/folders/1bxb--QlXr4oqN0DkXE6_5bWywhY3AQt0?usp=sharing", "review":"https://drive.google.com/drive/folders/1XHkysQiaxAevSWVrkgZaFvE8QPNUfMMl?usp=sharing"}, 
}

st.subheader("Subir informe")
st.markdown(f"[Abrir carpeta]({links[grupo]['upload']})")

st.subheader("Evaluar informes")
st.markdown(f"[Abrir informes]({links[grupo]['review']})")

# -----------------------------
# RÚBRICA
# -----------------------------
st.header("Evaluación entre grupos")

c1,c2 = st.columns(2)

with c1:
    coherencia = st.slider("Coherencia teórica",0,10,5)
    analisis = st.slider("Calidad del diagnóstico",0,10,5)

with c2:
    viabilidad = st.slider("Viabilidad",0,10,5)
    innovacion = st.slider("Innovación",0,10,5)

nota = (coherencia+analisis+viabilidad+innovacion)/4
st.metric("Nota final", round(nota,2))

comentario = st.text_area("Comentario")

# -----------------------------
# EXPORTAR EVALUACIÓN
# -----------------------------
if st.button("Generar informe de evaluación"):

    doc = Document()

    doc.add_heading('INFORME DE EVALUACIÓN',1)

    doc.add_paragraph(f"Grupo: {grupo}")
    doc.add_paragraph(f"Nota final: {round(nota,2)}")

    doc.add_heading("Puntuaciones",2)
    doc.add_paragraph(f"Coherencia: {coherencia}")
    doc.add_paragraph(f"Diagnóstico: {analisis}")
    doc.add_paragraph(f"Viabilidad: {viabilidad}")
    doc.add_paragraph(f"Innovación: {innovacion}")

    doc.add_heading("Comentario",2)
    doc.add_paragraph(comentario)

    buffer = BytesIO()
    doc.save(buffer)
    buffer.seek(0)

    st.download_button(
        "📥 Descargar evaluación",
        buffer,
        file_name="evaluacion.docx"
    )
